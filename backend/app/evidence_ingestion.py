import hashlib
import io
import json
import uuid
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import PurePosixPath

from docx import Document as DocxDocument
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.data_governance import create_deletion_request, execute_deletion_request
from app.data_governance_models import DeletionScope, DeletionStatus
from app.evidence_models import (
    EvidenceKind,
    EvidenceSource,
    EvidenceSourceStatus,
    EvidenceVisibility,
)
from app.models import (
    CanonicalEvent,
    IntegrationConnection,
    IntegrationHealth,
    IntegrationStatus,
    MembershipRole,
    RawEvent,
    RawEventStatus,
    ResourceAccessLevel,
    ResourceGrant,
)
from app.raw_events import persist_raw_event
from app.search_models import SearchDocument, SearchEmbeddingStatus
from app.security_audit import audit_authorization_decision
from app.data_governance import append_audit_event
from app.work_graph import project_canonical_event

GENERIC_EVIDENCE_PROVIDER = "generic_upload"
GENERIC_EVIDENCE_ACCOUNT = "brain:generic-upload-v1"
MAX_EVIDENCE_BYTES = 10_000_000
MAX_EXTRACTED_CHARS = 1_000_000
MAX_PDF_PAGES = 1_000
MAX_DOCX_UNCOMPRESSED_BYTES = 50_000_000
MAX_DOCX_ENTRIES = 10_000
CHUNK_CHARS = 6_000
CHUNK_OVERLAP_CHARS = 500
_TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".json", ".vtt", ".srt"}


class EvidenceIngestionError(ValueError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code[:128]


class EvidenceConflictError(EvidenceIngestionError):
    pass


@dataclass(frozen=True, slots=True)
class EvidenceChunk:
    index: int
    text: str
    sha256: str


def _safe_filename(value: str) -> str:
    normalized = value.replace("\\", "/")
    name = PurePosixPath(normalized).name.strip()
    if not name or name in {".", ".."}:
        raise EvidenceIngestionError("invalid_filename", "Evidence filename is required")
    return name[:512]


def _normalize_text(value: str) -> str:
    value = value.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    value = value.strip()
    if not value:
        raise EvidenceIngestionError("empty_extracted_text", "Evidence contains no extractable text")
    if len(value) > MAX_EXTRACTED_CHARS:
        raise EvidenceIngestionError(
            "extracted_text_too_large",
            "Extracted evidence text exceeds the supported limit",
        )
    return value


def _extract_plain_text(content: bytes) -> str:
    try:
        return _normalize_text(content.decode("utf-8-sig"))
    except UnicodeDecodeError as exc:
        raise EvidenceIngestionError(
            "invalid_text_encoding",
            "Text evidence must be UTF-8 encoded",
        ) from exc


def _extract_pdf(content: bytes) -> str:
    try:
        reader = PdfReader(io.BytesIO(content), strict=False)
    except Exception as exc:
        raise EvidenceIngestionError("invalid_pdf", "PDF could not be parsed") from exc
    if reader.is_encrypted:
        raise EvidenceIngestionError("encrypted_pdf", "Encrypted PDFs are not supported")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise EvidenceIngestionError("pdf_page_limit", "PDF exceeds the supported page limit")
    pages: list[str] = []
    try:
        for page_number, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"[Page {page_number}]\n{text}")
    except Exception as exc:
        raise EvidenceIngestionError("pdf_text_extraction_failed", "PDF text extraction failed") from exc
    return _normalize_text("\n\n".join(pages))


def _validate_docx_container(content: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_DOCX_ENTRIES:
                raise EvidenceIngestionError(
                    "docx_entry_limit",
                    "DOCX archive contains too many entries",
                )
            if sum(info.file_size for info in infos) > MAX_DOCX_UNCOMPRESSED_BYTES:
                raise EvidenceIngestionError(
                    "docx_expansion_limit",
                    "DOCX uncompressed size exceeds the supported limit",
                )
    except zipfile.BadZipFile as exc:
        raise EvidenceIngestionError("invalid_docx", "DOCX could not be parsed") from exc


def _extract_docx(content: bytes) -> str:
    _validate_docx_container(content)
    try:
        document = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise EvidenceIngestionError("invalid_docx", "DOCX could not be parsed") from exc
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if values:
                blocks.append(" | ".join(values))
    return _normalize_text("\n\n".join(blocks))


def extract_evidence_text(*, filename: str, media_type: str, content: bytes) -> str:
    suffix = PurePosixPath(filename.lower()).suffix
    normalized_media = media_type.split(";", 1)[0].strip().lower()
    if suffix == ".pdf" or normalized_media == "application/pdf":
        return _extract_pdf(content)
    if suffix == ".docx" or normalized_media == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    ):
        return _extract_docx(content)
    if (
        suffix in _TEXT_SUFFIXES
        or normalized_media.startswith("text/")
        or normalized_media in {"application/json", "application/csv"}
    ):
        return _extract_plain_text(content)
    raise EvidenceIngestionError(
        "unsupported_evidence_type",
        "Supported evidence types are UTF-8 text/Markdown/CSV/JSON/VTT/SRT, PDF and DOCX",
    )


def chunk_evidence_text(text: str) -> list[EvidenceChunk]:
    chunks: list[EvidenceChunk] = []
    start = 0
    index = 0
    while start < len(text):
        end = min(start + CHUNK_CHARS, len(text))
        if end < len(text):
            lower_bound = start + max(CHUNK_CHARS // 2, 1)
            newline = text.rfind("\n", lower_bound, end)
            space = text.rfind(" ", lower_bound, end)
            boundary = max(newline, space)
            if boundary > start:
                end = boundary
        chunk_text = text[start:end].strip()
        if chunk_text:
            chunks.append(
                EvidenceChunk(
                    index=index,
                    text=chunk_text,
                    sha256=hashlib.sha256(chunk_text.encode()).hexdigest(),
                )
            )
            index += 1
        if end >= len(text):
            break
        next_start = max(end - CHUNK_OVERLAP_CHARS, start + 1)
        start = next_start
    if not chunks:
        raise EvidenceIngestionError("empty_chunks", "Evidence produced no searchable chunks")
    return chunks


def _generic_connection(db: Session, organization_id: uuid.UUID, actor_user_id: uuid.UUID):
    connection = db.scalar(
        select(IntegrationConnection).where(
            IntegrationConnection.organization_id == organization_id,
            IntegrationConnection.provider == GENERIC_EVIDENCE_PROVIDER,
            IntegrationConnection.external_account_id == GENERIC_EVIDENCE_ACCOUNT,
        )
    )
    if connection is not None:
        if connection.status != IntegrationStatus.ACTIVE:
            raise EvidenceConflictError(
                "generic_adapter_disabled",
                "Generic evidence ingestion is disabled for this organization",
            )
        return connection

    connection = IntegrationConnection(
        organization_id=organization_id,
        provider=GENERIC_EVIDENCE_PROVIDER,
        external_account_id=GENERIC_EVIDENCE_ACCOUNT,
        display_name="Brain generic file/transcript upload",
        status=IntegrationStatus.ACTIVE,
        health=IntegrationHealth.HEALTHY,
        scopes=["evidence:upload"],
        provider_metadata={"managed_by": "brain", "adapter": "generic-upload-v1"},
        secret_ref=None,
        created_by_user_id=actor_user_id,
    )
    db.add(connection)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        connection = db.scalar(
            select(IntegrationConnection).where(
                IntegrationConnection.organization_id == organization_id,
                IntegrationConnection.provider == GENERIC_EVIDENCE_PROVIDER,
                IntegrationConnection.external_account_id == GENERIC_EVIDENCE_ACCOUNT,
            )
        )
        if connection is None or connection.status != IntegrationStatus.ACTIVE:
            raise EvidenceConflictError(
                "generic_adapter_unavailable",
                "Generic evidence ingestion connection could not be created",
            ) from None
    db.refresh(connection)
    return connection


def _existing_idempotent_source(
    db: Session,
    *,
    organization_id: uuid.UUID,
    idempotency_key: str | None,
    content_sha256: str,
    kind: EvidenceKind,
) -> EvidenceSource | None:
    if idempotency_key is None:
        return None
    existing = db.scalar(
        select(EvidenceSource).where(
            EvidenceSource.organization_id == organization_id,
            EvidenceSource.idempotency_key == idempotency_key,
        )
    )
    if existing is None:
        return None
    if (
        existing.content_sha256 != content_sha256
        or existing.kind != kind
        or existing.status != EvidenceSourceStatus.ACTIVE
    ):
        raise EvidenceConflictError(
            "idempotency_key_reused",
            "Idempotency key was already used for different or non-active evidence",
        )
    return existing


def _chunk_payload(
    source: EvidenceSource,
    chunk: EvidenceChunk,
    *,
    chunk_count: int,
    actor_user_id: uuid.UUID,
) -> bytes:
    payload = {
        "evidence_source_id": str(source.id),
        "kind": source.kind.value,
        "title": source.title,
        "filename": source.filename,
        "media_type": source.media_type,
        "source_sha256": source.content_sha256,
        "chunk_index": chunk.index,
        "chunk_count": chunk_count,
        "chunk_sha256": chunk.sha256,
        "text": chunk.text,
        "uploaded_by_user_id": str(actor_user_id),
        "occurred_at": source.occurred_at.isoformat() if source.occurred_at else None,
    }
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":")).encode()


def _canonical_chunk(
    db: Session,
    *,
    source: EvidenceSource,
    raw: RawEvent,
    chunk: EvidenceChunk,
    chunk_count: int,
    actor_user_id: uuid.UUID,
) -> tuple[CanonicalEvent, uuid.UUID]:
    existing = db.scalar(
        select(CanonicalEvent).where(CanonicalEvent.raw_event_id == raw.id)
    )
    if existing is None:
        occurred_at = source.occurred_at or datetime.now(UTC)
        provenance = {
            "raw_event_id": str(raw.id),
            "integration_connection_id": str(source.integration_connection_id),
            "payload_sha256": raw.payload_sha256,
            "delivery_kind": raw.delivery_kind,
            "evidence_source_id": str(source.id),
            "source_sha256": source.content_sha256,
            "chunk_sha256": chunk.sha256,
            "chunk_index": chunk.index,
            "chunk_count": chunk_count,
            "filename": source.filename,
            "media_type": source.media_type,
        }
        existing = CanonicalEvent(
            organization_id=source.organization_id,
            raw_event_id=raw.id,
            integration_connection_id=source.integration_connection_id,
            resolved_user_id=actor_user_id,
            schema_version=1,
            event_type=f"{source.kind.value}.observed",
            action="observed",
            actor_type="brain_user",
            actor_external_id=str(actor_user_id),
            actor_display_name=None,
            object_type=source.kind.value,
            object_external_id=str(source.id),
            object_display_name=source.title,
            source_provider=GENERIC_EVIDENCE_PROVIDER,
            source_event_id=raw.source_event_id,
            source_event_type=raw.source_event_type,
            occurred_at=occurred_at,
            source_visibility=source.source_visibility.value,
            source_acl=list(source.source_acl),
            provenance=provenance,
            event_metadata={**provenance, "text": chunk.text},
        )
        raw.processing_status = RawEventStatus.PROCESSED
        raw.last_error_code = None
        db.add(existing)
        db.commit()
        db.refresh(existing)
    node = project_canonical_event(db, existing)
    return existing, node.id


def _upsert_search_chunk(
    db: Session,
    *,
    source: EvidenceSource,
    event: CanonicalEvent,
    work_graph_node_id: uuid.UUID,
    chunk: EvidenceChunk,
    chunk_count: int,
) -> SearchDocument:
    document = db.scalar(
        select(SearchDocument).where(SearchDocument.canonical_event_id == event.id)
    )
    provenance = {
        **event.provenance,
        "source_event_id": event.source_event_id,
        "source_event_type": event.source_event_type,
        "object_external_id": event.object_external_id,
        "chunk_index": chunk.index,
        "chunk_count": chunk_count,
    }
    if document is None:
        document = SearchDocument(
            organization_id=source.organization_id,
            canonical_event_id=event.id,
            integration_connection_id=source.integration_connection_id,
            work_graph_node_id=work_graph_node_id,
            source_provider=GENERIC_EVIDENCE_PROVIDER,
            source_visibility=source.source_visibility.value,
            source_acl=list(source.source_acl),
            channel_id=None,
            repository_id=None,
            object_type=source.kind.value,
            object_external_id=str(source.id),
            title=source.title,
            content=chunk.text,
            provenance=provenance,
            occurred_at=event.occurred_at,
            is_deleted=False,
            embedding_status=SearchEmbeddingStatus.PENDING,
        )
        db.add(document)
    else:
        document.work_graph_node_id = work_graph_node_id
        document.source_visibility = source.source_visibility.value
        document.source_acl = list(source.source_acl)
        document.title = source.title
        document.content = chunk.text
        document.provenance = provenance
        document.occurred_at = event.occurred_at
        document.is_deleted = False
        document.embedding = None
        document.embedding_model = None
        document.embedding_status = SearchEmbeddingStatus.PENDING
        document.embedding_attempts = 0
        document.next_retry_at = None
        document.claimed_at = None
        document.last_error_code = None
    return document


def _grant_restricted_chunk(
    db: Session,
    *,
    organization_id: uuid.UUID,
    actor_user_id: uuid.UUID,
    work_graph_node_id: uuid.UUID,
) -> None:
    resource_id = str(work_graph_node_id)
    existing = db.scalar(
        select(ResourceGrant.id).where(
            ResourceGrant.organization_id == organization_id,
            ResourceGrant.resource_type == "work_graph.node",
            ResourceGrant.resource_id == resource_id,
            ResourceGrant.user_id == actor_user_id,
            ResourceGrant.access == ResourceAccessLevel.WRITE,
        )
    )
    if existing is None:
        db.add(
            ResourceGrant(
                organization_id=organization_id,
                resource_type="work_graph.node",
                resource_id=resource_id,
                user_id=actor_user_id,
                access=ResourceAccessLevel.WRITE,
                created_by_user_id=actor_user_id,
            )
        )


def ingest_evidence(
    db: Session,
    *,
    organization_id: uuid.UUID,
    actor_user_id: uuid.UUID,
    kind: EvidenceKind,
    title: str | None,
    filename: str,
    media_type: str,
    content: bytes,
    visibility: EvidenceVisibility,
    occurred_at: datetime | None,
    idempotency_key: str | None,
    request_id: str | None = None,
) -> EvidenceSource:
    if not content:
        raise EvidenceIngestionError("empty_upload", "Evidence upload must not be empty")
    if len(content) > MAX_EVIDENCE_BYTES:
        raise EvidenceIngestionError("upload_too_large", "Evidence upload exceeds 10 MB")
    safe_filename = _safe_filename(filename)
    normalized_title = " ".join((title or safe_filename).strip().split())[:512]
    if not normalized_title:
        raise EvidenceIngestionError("invalid_title", "Evidence title is required")
    normalized_media = (media_type or "application/octet-stream").strip()[:128]
    normalized_key = " ".join(idempotency_key.strip().split())[:128] if idempotency_key else None
    if idempotency_key and not normalized_key:
        raise EvidenceIngestionError("invalid_idempotency_key", "Idempotency key is invalid")

    content_sha256 = hashlib.sha256(content).hexdigest()
    existing = _existing_idempotent_source(
        db,
        organization_id=organization_id,
        idempotency_key=normalized_key,
        content_sha256=content_sha256,
        kind=kind,
    )
    if existing is not None:
        return existing

    extracted = extract_evidence_text(
        filename=safe_filename,
        media_type=normalized_media,
        content=content,
    )
    chunks = chunk_evidence_text(extracted)
    connection = _generic_connection(db, organization_id, actor_user_id)
    source_acl = [str(actor_user_id)] if visibility == EvidenceVisibility.RESTRICTED else []
    source = EvidenceSource(
        organization_id=organization_id,
        integration_connection_id=connection.id,
        kind=kind,
        title=normalized_title,
        filename=safe_filename,
        media_type=normalized_media,
        content_sha256=content_sha256,
        byte_size=len(content),
        raw_content=content,
        source_visibility=visibility,
        source_acl=source_acl,
        chunk_count=0,
        extracted_char_count=0,
        idempotency_key=normalized_key,
        created_by_user_id=actor_user_id,
        occurred_at=occurred_at,
        status=EvidenceSourceStatus.PROCESSING,
    )
    db.add(source)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        replay = _existing_idempotent_source(
            db,
            organization_id=organization_id,
            idempotency_key=normalized_key,
            content_sha256=content_sha256,
            kind=kind,
        )
        if replay is not None:
            return replay
        raise EvidenceConflictError(
            "evidence_source_conflict",
            "Evidence source could not be created because of a conflicting upload",
        ) from exc
    db.refresh(source)

    projected: list[tuple[CanonicalEvent, uuid.UUID, EvidenceChunk]] = []
    try:
        for chunk in chunks:
            raw_result = persist_raw_event(
                db,
                organization_id=organization_id,
                integration_connection_id=connection.id,
                provider=GENERIC_EVIDENCE_PROVIDER,
                source_event_id=f"{source.id}:chunk:{chunk.index}:{chunk.sha256[:16]}",
                source_event_type="evidence.chunk",
                delivery_kind="upload",
                raw_payload=_chunk_payload(
                    source,
                    chunk,
                    chunk_count=len(chunks),
                    actor_user_id=actor_user_id,
                ),
                content_type="application/json",
                source_visibility=visibility.value,
                source_acl=source_acl,
                source_timestamp=occurred_at,
            )
            event, node_id = _canonical_chunk(
                db,
                source=source,
                raw=raw_result.event,
                chunk=chunk,
                chunk_count=len(chunks),
                actor_user_id=actor_user_id,
            )
            projected.append((event, node_id, chunk))

        for event, node_id, chunk in projected:
            _upsert_search_chunk(
                db,
                source=source,
                event=event,
                work_graph_node_id=node_id,
                chunk=chunk,
                chunk_count=len(chunks),
            )
            if visibility == EvidenceVisibility.RESTRICTED:
                _grant_restricted_chunk(
                    db,
                    organization_id=organization_id,
                    actor_user_id=actor_user_id,
                    work_graph_node_id=node_id,
                )
        source.chunk_count = len(chunks)
        source.extracted_char_count = len(extracted)
        source.status = EvidenceSourceStatus.ACTIVE
        source.last_error_code = None
        db.commit()
        db.refresh(source)
        append_audit_event(
            db,
            organization_id=organization_id,
            event_key=f"evidence.uploaded:{source.id}",
            event_type="evidence.uploaded",
            outcome="succeeded",
            actor_user_id=actor_user_id,
            resource_type="evidence_source",
            resource_id=source.id,
            request_id=request_id,
            metadata={
                "kind": kind.value,
                "filename": safe_filename,
                "media_type": normalized_media,
                "source_sha256": content_sha256,
                "byte_size": len(content),
                "chunk_count": len(chunks),
                "visibility": visibility.value,
            },
        )
        return source
    except Exception as exc:
        db.rollback()
        persisted = db.get(EvidenceSource, source.id)
        if persisted is not None and persisted.status == EvidenceSourceStatus.PROCESSING:
            persisted.status = EvidenceSourceStatus.FAILED
            persisted.last_error_code = getattr(exc, "code", "evidence_projection_failed")[:128]
            db.commit()
        raise


def evidence_source_visible_to_user(source: EvidenceSource, user_id: uuid.UUID) -> bool:
    return (
        source.source_visibility == EvidenceVisibility.ORGANIZATION
        or str(user_id) in source.source_acl
    )


def get_visible_evidence_source(
    db: Session,
    *,
    organization_id: uuid.UUID,
    user_id: uuid.UUID,
    source_id: uuid.UUID,
) -> EvidenceSource | None:
    source = db.scalar(
        select(EvidenceSource).where(
            EvidenceSource.id == source_id,
            EvidenceSource.organization_id == organization_id,
        )
    )
    if source is None or not evidence_source_visible_to_user(source, user_id):
        return None
    return source


def list_visible_evidence_sources(
    db: Session,
    *,
    organization_id: uuid.UUID,
    user_id: uuid.UUID,
    status: EvidenceSourceStatus | None,
    limit: int,
) -> list[EvidenceSource]:
    query = select(EvidenceSource).where(EvidenceSource.organization_id == organization_id)
    if status is not None:
        query = query.where(EvidenceSource.status == status)
    rows = list(
        db.scalars(
            query.order_by(EvidenceSource.created_at.desc(), EvidenceSource.id.desc()).limit(limit)
        )
    )
    return [row for row in rows if evidence_source_visible_to_user(row, user_id)]


def delete_evidence_source(
    db: Session,
    *,
    organization_id: uuid.UUID,
    actor_user_id: uuid.UUID,
    actor_role: MembershipRole,
    source_id: uuid.UUID,
    request_id: str | None = None,
) -> EvidenceSource:
    source = db.scalar(
        select(EvidenceSource).where(
            EvidenceSource.id == source_id,
            EvidenceSource.organization_id == organization_id,
        )
    )
    if source is None:
        raise EvidenceIngestionError("evidence_not_found", "Evidence source not found")
    if source.status == EvidenceSourceStatus.DELETED:
        return source
    can_delete = source.created_by_user_id == actor_user_id or actor_role in {
        MembershipRole.OWNER,
        MembershipRole.ADMIN,
    }
    if not can_delete:
        audit_authorization_decision(
            allowed=False,
            organization_id=organization_id,
            actor_user_id=actor_user_id,
            permission="evidence.delete",
            reason="not_source_owner",
            resource_type="evidence_source",
            resource_id=str(source.id),
            db=db,
        )
        raise EvidenceIngestionError("evidence_not_found", "Evidence source not found")

    deletion = create_deletion_request(
        db,
        organization_id=organization_id,
        actor_user_id=actor_user_id,
        request_key=f"generic-evidence-delete:{source.id}",
        scope=DeletionScope.SOURCE_OBJECT,
        reason="Evidence source deleted by authorized Brain user",
        integration_connection_id=source.integration_connection_id,
        source_provider=GENERIC_EVIDENCE_PROVIDER,
        object_type=source.kind.value,
        object_external_id=str(source.id),
        request_id=request_id,
    )
    deletion = execute_deletion_request(
        db,
        organization_id=organization_id,
        deletion_request_id=deletion.id,
        request_id=request_id,
    )
    if deletion.status != DeletionStatus.COMPLETED:
        raise EvidenceConflictError(
            "evidence_delete_incomplete",
            "Evidence deletion did not complete",
        )
    source = db.get(EvidenceSource, source_id)
    if source is None:
        raise EvidenceIngestionError("evidence_not_found", "Evidence source not found")
    source.raw_content = None
    source.status = EvidenceSourceStatus.DELETED
    source.deleted_at = datetime.now(UTC)
    source.last_error_code = None
    db.commit()
    db.refresh(source)
    append_audit_event(
        db,
        organization_id=organization_id,
        event_key=f"evidence.deleted:{source.id}",
        event_type="evidence.deleted",
        outcome="succeeded",
        actor_user_id=actor_user_id,
        resource_type="evidence_source",
        resource_id=source.id,
        request_id=request_id,
        metadata={
            "kind": source.kind.value,
            "source_sha256": source.content_sha256,
            "raw_content_removed": True,
        },
    )
    return source
