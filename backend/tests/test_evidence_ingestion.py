import io
import uuid

from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.auth import get_current_user
from app.evidence_ingestion import extract_evidence_text
from app.evidence_models import EvidenceSource, EvidenceSourceStatus
from app.main import app
from app.models import (
    CanonicalEvent,
    IntegrationConnection,
    IntegrationStatus,
    Membership,
    MembershipRole,
    Organization,
    User,
)
from app.search import SearchMode, search_documents
from app.search_models import SearchDocument


def _seed(db: Session, suffix: str):
    owner = User(email=f"evidence-owner-{suffix}@example.com")
    member = User(email=f"evidence-member-{suffix}@example.com")
    other = User(email=f"evidence-other-{suffix}@example.com")
    guest = User(email=f"evidence-guest-{suffix}@example.com")
    organization = Organization(name=f"Evidence {suffix}", slug=f"evidence-{suffix}")
    db.add_all([owner, member, other, guest, organization])
    db.flush()
    db.add_all(
        [
            Membership(
                organization_id=organization.id,
                user_id=owner.id,
                role=MembershipRole.OWNER,
            ),
            Membership(
                organization_id=organization.id,
                user_id=member.id,
                role=MembershipRole.MEMBER,
            ),
            Membership(
                organization_id=organization.id,
                user_id=other.id,
                role=MembershipRole.MEMBER,
            ),
            Membership(
                organization_id=organization.id,
                user_id=guest.id,
                role=MembershipRole.GUEST,
            ),
        ]
    )
    db.commit()
    return organization, owner, member, other, guest


def _upload(client, organization, user, *, visibility="organization", key="evidence-key"):
    app.dependency_overrides[get_current_user] = lambda: user
    return client.post(
        f"/api/v1/organizations/{organization.id}/evidence/uploads",
        headers={"Idempotency-Key": key},
        data={
            "kind": "document",
            "title": "Launch notes",
            "visibility": visibility,
        },
        files={
            "file": (
                "launch-notes.txt",
                b"Project Atlas launch decision: ship the guarded release on Friday.",
                "text/plain",
            )
        },
    )


def test_upload_projects_immutable_provenance_into_search(
    db_session: Session,
    client,
) -> None:
    organization, _, member, _, _ = _seed(db_session, "projection")
    response = _upload(client, organization, member)
    assert response.status_code == 201
    source_id = response.json()["id"]

    source = db_session.get(EvidenceSource, uuid.UUID(source_id))
    assert source is not None
    assert source.status == EvidenceSourceStatus.ACTIVE
    assert source.raw_content is not None
    assert len(source.content_sha256) == 64
    assert source.chunk_count >= 1

    documents = list(
        db_session.scalars(
            select(SearchDocument).where(
                SearchDocument.organization_id == organization.id,
                SearchDocument.object_external_id == source_id,
            )
        )
    )
    assert documents
    assert all(document.source_provider == "generic_upload" for document in documents)
    assert all(
        document.provenance["source_sha256"] == source.content_sha256
        for document in documents
    )
    assert all(document.provenance["evidence_source_id"] == source_id for document in documents)

    canonical = list(
        db_session.scalars(
            select(CanonicalEvent).where(
                CanonicalEvent.organization_id == organization.id,
                CanonicalEvent.object_external_id == source_id,
            )
        )
    )
    assert canonical
    assert all(event.source_provider == "generic_upload" for event in canonical)

    result = search_documents(
        db_session,
        organization_id=organization.id,
        user_id=member.id,
        query="Atlas launch",
        mode=SearchMode.KEYWORD,
        limit=10,
        embedding_client=None,
        embedding_model=None,
    )
    assert any(hit.document.object_external_id == source_id for hit in result.hits)


def test_upload_idempotency_returns_same_source_and_rejects_key_reuse(
    db_session: Session,
    client,
) -> None:
    organization, _, member, _, _ = _seed(db_session, "idempotency")
    first = _upload(client, organization, member, key="same-key")
    second = _upload(client, organization, member, key="same-key")
    assert first.status_code == 201
    assert second.status_code == 201
    assert second.json()["id"] == first.json()["id"]

    different = client.post(
        f"/api/v1/organizations/{organization.id}/evidence/uploads",
        headers={"Idempotency-Key": "same-key"},
        data={"kind": "document", "visibility": "organization"},
        files={"file": ("other.txt", b"different evidence", "text/plain")},
    )
    assert different.status_code == 409
    assert different.json()["detail"]["code"] == "idempotency_key_reused"


def test_restricted_evidence_is_searchable_only_by_granted_uploader(
    db_session: Session,
    client,
) -> None:
    organization, _, member, other, _ = _seed(db_session, "restricted")
    uploaded = _upload(
        client,
        organization,
        member,
        visibility="restricted",
        key="restricted-key",
    )
    assert uploaded.status_code == 201
    source_id = uploaded.json()["id"]

    visible = search_documents(
        db_session,
        organization_id=organization.id,
        user_id=member.id,
        query="Atlas launch",
        mode=SearchMode.KEYWORD,
        limit=10,
        embedding_client=None,
        embedding_model=None,
    )
    hidden = search_documents(
        db_session,
        organization_id=organization.id,
        user_id=other.id,
        query="Atlas launch",
        mode=SearchMode.KEYWORD,
        limit=10,
        embedding_client=None,
        embedding_model=None,
    )
    assert any(hit.document.object_external_id == source_id for hit in visible.hits)
    assert all(hit.document.object_external_id != source_id for hit in hidden.hits)

    app.dependency_overrides[get_current_user] = lambda: other
    metadata = client.get(f"/api/v1/organizations/{organization.id}/evidence/{source_id}")
    assert metadata.status_code == 404


def test_delete_physically_removes_source_bytes_and_searchable_derivatives(
    db_session: Session,
    client,
) -> None:
    organization, _, member, _, _ = _seed(db_session, "delete")
    uploaded = _upload(client, organization, member, key="delete-key")
    assert uploaded.status_code == 201
    source_id = uploaded.json()["id"]

    deleted = client.delete(f"/api/v1/organizations/{organization.id}/evidence/{source_id}")
    assert deleted.status_code == 200
    source = db_session.get(EvidenceSource, uuid.UUID(source_id))
    assert source is not None
    assert source.status == EvidenceSourceStatus.DELETED
    assert source.raw_content is None
    assert list(
        db_session.scalars(
            select(SearchDocument).where(SearchDocument.object_external_id == source_id)
        )
    ) == []
    assert list(
        db_session.scalars(
            select(CanonicalEvent).where(CanonicalEvent.object_external_id == source_id)
        )
    ) == []


def test_revoked_generic_adapter_disappears_from_search(
    db_session: Session,
    client,
) -> None:
    organization, _, member, _, _ = _seed(db_session, "revoke")
    uploaded = _upload(client, organization, member, key="revoke-key")
    assert uploaded.status_code == 201
    source_id = uploaded.json()["id"]
    source = db_session.get(EvidenceSource, uuid.UUID(source_id))
    assert source is not None
    connection = db_session.get(IntegrationConnection, source.integration_connection_id)
    assert connection is not None
    connection.status = IntegrationStatus.REVOKED
    db_session.commit()

    result = search_documents(
        db_session,
        organization_id=organization.id,
        user_id=member.id,
        query="Atlas launch",
        mode=SearchMode.KEYWORD,
        limit=10,
        embedding_client=None,
        embedding_model=None,
    )
    assert all(hit.document.object_external_id != source_id for hit in result.hits)


def test_guest_cannot_upload_and_unsupported_binary_is_rejected(
    db_session: Session,
    client,
) -> None:
    organization, _, member, _, guest = _seed(db_session, "negative")
    app.dependency_overrides[get_current_user] = lambda: guest
    denied = client.post(
        f"/api/v1/organizations/{organization.id}/evidence/uploads",
        data={"kind": "document", "visibility": "organization"},
        files={"file": ("notes.txt", b"content", "text/plain")},
    )
    assert denied.status_code == 403

    app.dependency_overrides[get_current_user] = lambda: member
    unsupported = client.post(
        f"/api/v1/organizations/{organization.id}/evidence/uploads",
        data={"kind": "document", "visibility": "organization"},
        files={"file": ("archive.bin", b"\x00\x01\x02", "application/octet-stream")},
    )
    assert unsupported.status_code == 400
    assert unsupported.json()["detail"]["code"] == "unsupported_evidence_type"


def test_docx_extractor_reads_paragraphs_and_table_cells() -> None:
    document = DocxDocument()
    document.add_paragraph("Decision: use evidence-backed project reporting.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Platform"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_evidence_text(
        filename="meeting.docx",
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=buffer.getvalue(),
    )
    assert "evidence-backed project reporting" in text
    assert "Owner | Platform" in text
